"""
Convert Markdown to Word (.docx) with linked Table of Contents
and outline formatting.

Usage:
    python3 scripts/md_to_docx.py <input.md> <output.docx>

This is not a pre-commit hook and it needs third-party `python-docx`, so it must be run against a
real interpreter rather than through `pre-commit`. The name is OS-dependent — `python3` on
macOS/Linux, `python` on Windows (docs/orchestration/knowledge/reference_precommit_interpreter.md).
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import copy


def setup_styles(doc):
    """Configure document styles for professional formatting."""
    style = doc.styles

    # Normal style
    normal = style['Normal']
    normal.font.name = 'Catamaran'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    # Title style
    title = style['Title']
    title.font.name = 'Titillium Web'
    title.font.size = Pt(28)
    title.font.color.rgb = RGBColor(0, 51, 102)  # Dark navy
    title.font.bold = True
    title.paragraph_format.space_after = Pt(4)

    # Subtitle style
    subtitle = style['Subtitle']
    subtitle.font.name = 'Catamaran'
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = RGBColor(89, 89, 89)
    subtitle.font.italic = True
    subtitle.paragraph_format.space_after = Pt(12)

    # Heading 1
    h1 = style['Heading 1']
    h1.font.name = 'Titillium Web'
    h1.font.size = Pt(20)
    h1.font.color.rgb = RGBColor(0, 51, 102)
    h1.font.bold = True
    h1.paragraph_format.space_before = Pt(24)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.page_break_before = True

    # Heading 2
    h2 = style['Heading 2']
    h2.font.name = 'Titillium Web'
    h2.font.size = Pt(16)
    h2.font.color.rgb = RGBColor(0, 76, 153)
    h2.font.bold = True
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.page_break_before = False

    # Heading 3
    h3 = style['Heading 3']
    h3.font.name = 'Catamaran'
    h3.font.size = Pt(13)
    h3.font.color.rgb = RGBColor(0, 76, 153)
    h3.font.bold = True
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)

    # Create a style for code blocks
    if 'Code Block' not in [s.name for s in style]:
        code_style = style.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = 'Menlo'
        code_style.font.size = Pt(8.5)
        code_style.font.color.rgb = RGBColor(40, 40, 40)
        code_style.paragraph_format.space_before = Pt(2)
        code_style.paragraph_format.space_after = Pt(2)
        code_style.paragraph_format.line_spacing = 1.0
        code_style.paragraph_format.left_indent = Inches(0.25)
        # Add shading
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F5F5"/>')
        code_style.element.get_or_add_pPr().append(shading)

    # Create a style for metadata lines (Prepared, Classification, etc.)
    if 'Metadata' not in [s.name for s in style]:
        meta_style = style.add_style('Metadata', WD_STYLE_TYPE.PARAGRAPH)
        meta_style.font.name = 'Catamaran'
        meta_style.font.size = Pt(10)
        meta_style.font.color.rgb = RGBColor(89, 89, 89)
        meta_style.paragraph_format.space_after = Pt(2)

    return style


def add_toc(doc):
    """Add a linked Table of Contents field."""
    para = doc.add_paragraph()
    para.style = doc.styles['Heading 1']
    para.style.paragraph_format.page_break_before = False  # No page break for TOC heading
    run = para.add_run('Table of Contents')

    # Add a TOC field
    para2 = doc.add_paragraph()
    run2 = para2.add_run()
    fld_char_begin = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run2._r.append(fld_char_begin)

    run3 = para2.add_run()
    instr_text = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>')
    run3._r.append(instr_text)

    run4 = para2.add_run()
    fld_char_separate = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    run4._r.append(fld_char_separate)

    run5 = para2.add_run('[Right-click and select "Update Field" to populate Table of Contents]')
    run5.font.color.rgb = RGBColor(128, 128, 128)
    run5.font.italic = True

    run6 = para2.add_run()
    fld_char_end = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run6._r.append(fld_char_end)

    # Add a page break after the TOC
    doc.add_page_break()


def add_formatted_run(paragraph, text):
    """Add a run with inline bold/italic markdown formatting."""
    # Split on bold markers first
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            inner = part[2:-2]
            # Check for italic inside bold
            italic_parts = re.split(r'(\*.*?\*)', inner)
            for ip in italic_parts:
                if ip.startswith('*') and ip.endswith('*'):
                    run = paragraph.add_run(ip[1:-1])
                    run.bold = True
                    run.italic = True
                else:
                    run = paragraph.add_run(ip)
                    run.bold = True
        else:
            # Check for standalone italic
            italic_parts = re.split(r'(\*[^*]+?\*)', part)
            for ip in italic_parts:
                if ip.startswith('*') and ip.endswith('*') and not ip.startswith('**'):
                    run = paragraph.add_run(ip[1:-1])
                    run.italic = True
                else:
                    paragraph.add_run(ip)


def add_bullet(doc, text, level=0):
    """Add a bullet point with proper formatting."""
    para = doc.add_paragraph(style='List Bullet')
    para.paragraph_format.left_indent = Inches(0.25 + level * 0.25)
    para.paragraph_format.space_after = Pt(4)
    add_formatted_run(para, text)
    return para


def add_table(doc, headers, rows):
    """Add a formatted table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Set header row
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        para = cell.paragraphs[0]
        run = para.add_run(header.strip().replace('**', ''))
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = 'Catamaran'
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # Header shading
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="003366" w:themeFill="text2"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Set data rows
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            para = cell.paragraphs[0]
            cleaned = cell_text.strip()
            add_formatted_run(para, cleaned)
            for run in para.runs:
                run.font.size = Pt(10)
                run.font.name = 'Catamaran'
            # Alternate row shading
            if r_idx % 2 == 1:
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8F0FE"/>')
                cell._tc.get_or_add_tcPr().append(shading)

    doc.add_paragraph()  # Spacing after table
    return table


def parse_markdown_table(lines, start_idx):
    """Parse a markdown table starting at start_idx. Returns (headers, rows, end_idx)."""
    headers = [c.strip() for c in lines[start_idx].split('|')[1:-1]]
    # Skip separator line
    row_idx = start_idx + 2
    rows = []
    while row_idx < len(lines) and '|' in lines[row_idx] and lines[row_idx].strip().startswith('|'):
        cells = [c.strip() for c in lines[row_idx].split('|')[1:-1]]
        # Pad or trim to match header count
        while len(cells) < len(headers):
            cells.append('')
        rows.append(cells[:len(headers)])
        row_idx += 1
    return headers, rows, row_idx


def convert_md_to_docx(md_path, docx_path):
    """Main conversion function."""
    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.page_height = Inches(11)
    section.page_width = Inches(8.5)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    setup_styles(doc)

    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')

    # Track state
    in_code_block = False
    code_lines = []
    i = 0
    first_h1_seen = False
    toc_added = False

    while i < len(lines):
        line = lines[i]

        # Code block handling
        if line.strip().startswith('```'):
            if in_code_block:
                # End code block - write accumulated lines
                for cl in code_lines:
                    p = doc.add_paragraph(cl, style='Code Block')
                code_lines = []
                in_code_block = False
                i += 1
                continue
            else:
                in_code_block = True
                code_lines = []
                i += 1
                continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Skip horizontal rules
        if line.strip() == '---':
            i += 1
            continue

        # Empty line
        if not line.strip():
            i += 1
            continue

        # Title (# heading)
        if line.startswith('# ') and not line.startswith('## '):
            title_text = line[2:].strip()
            p = doc.add_paragraph(title_text, style='Title')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Subtitle line
        if line.startswith('## ') and 'Executive Summary' in line:
            subtitle_text = line[3:].strip()
            p = doc.add_paragraph(subtitle_text, style='Subtitle')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Metadata lines (bold key: value)
        if line.startswith('**') and ':**' in line and not first_h1_seen:
            cleaned = line.replace('**', '')
            p = doc.add_paragraph(cleaned, style='Metadata')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Heading 1 (## )
        if line.startswith('## ') and not line.startswith('### '):
            heading_text = line[3:].strip()

            # Add TOC before the first real content heading
            if not toc_added:
                add_toc(doc)
                toc_added = True
                # Re-enable page break for H1 after TOC
                doc.styles['Heading 1'].paragraph_format.page_break_before = True

            first_h1_seen = True
            p = doc.add_paragraph(heading_text, style='Heading 1')
            i += 1
            continue

        # Heading 2 (### )
        if line.startswith('### ') and not line.startswith('#### '):
            heading_text = line[4:].strip()
            p = doc.add_paragraph(heading_text, style='Heading 2')
            i += 1
            continue

        # Heading 3 (#### )
        if line.startswith('#### '):
            heading_text = line[5:].strip()
            p = doc.add_paragraph(heading_text, style='Heading 3')
            i += 1
            continue

        # Table
        if '|' in line and line.strip().startswith('|'):
            # Check if next line is separator
            if i + 1 < len(lines) and re.match(r'\|[\s\-:|]+\|', lines[i + 1]):
                headers, rows, end_idx = parse_markdown_table(lines, i)
                add_table(doc, headers, rows)
                i = end_idx
                continue

        # Bullet point
        if line.strip().startswith('- '):
            text = line.strip()[2:].strip()
            add_bullet(doc, text)
            i += 1
            continue

        # Regular paragraph
        text = line.strip()
        if text:
            p = doc.add_paragraph()
            add_formatted_run(p, text)

        i += 1

    # Add footer with document source info
    footer_section = doc.sections[0]
    footer = footer_section.footer
    footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_para.add_run('Generated from Markdown — Draft')
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.font.name = 'Catamaran'

    # Add page numbers to footer
    page_num_para = footer.add_paragraph()
    page_num_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = page_num_para.add_run()
    fld_char_begin = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run._r.append(fld_char_begin)
    run2 = page_num_para.add_run()
    instr = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    run2._r.append(instr)
    run3 = page_num_para.add_run()
    fld_char_end = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run3._r.append(fld_char_end)

    doc.save(str(docx_path))
    print(f"Word document saved to: {docx_path}")


if __name__ == '__main__':
    import sys
    if len(sys.argv) != 3:
        # Name the interpreter that actually launched us, so the hint is right on every OS.
        print(f"Usage: {Path(sys.executable).name} scripts/md_to_docx.py <input.md> <output.docx>")
        sys.exit(1)
    md_file = Path(sys.argv[1])
    docx_file = Path(sys.argv[2])
    convert_md_to_docx(md_file, docx_file)
